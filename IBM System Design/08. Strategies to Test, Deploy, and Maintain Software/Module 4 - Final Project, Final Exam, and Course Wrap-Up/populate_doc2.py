import docx
import os

dir_path = r'd:\Ent\Learning\Cousera\IBM System Design\08. Strategies to Test, Deploy, and Maintain Software\Module 4 - Final Project, Final Exam, and Course Wrap-Up'
src2 = os.path.join(dir_path, 'NFR-template.docx')
out2 = os.path.join(dir_path, 'NFR-Completed.docx')

doc = docx.Document(src2)

# Update metadata
for p in doc.paragraphs:
    if '[Insert your full name]' in p.text:
        p.text = p.text.replace('[Insert your full name]', 'Systems Professional (NexaCare Solutions)')
    if '[Insert date]' in p.text:
        p.text = p.text.replace('[Insert date]', 'October 19, 2026')

# Table 1: NFR Definition (Table 0 in doc2 is often a small layout table, let's inspect)
# Let's find tables that have matching headers
t_nfr = None
t_golden = None
t_peak = None

for t in doc.tables:
    header = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]
    if 'NFR Category' in header:
        t_nfr = t
    elif 'Golden Signal' in header:
        t_golden = t
    elif 'Normal Threshold' in header:
        t_peak = t

# Populate NFR Table
if t_nfr:
    nfr_data = [
        ("Performance", "Appointment booking transaction completion time", "< 3.0 seconds", "End-to-end client telemetry and APM transaction tracing"),
        ("Performance", "Doctor availability update propagation across sessions", "< 5.0 seconds", "WebSocket broadcast delivery latency measurement"),
        ("Performance", "SMS reminder notification dispatch latency", "< 2.0 minutes from schedule", "Scheduled queue worker delivery receipt (DLR) timestamp"),
        ("Scalability", "Baseline concurrent active users", "500 users", "Sustained load testing validation via JMeter/k6"),
        ("Scalability", "Monday morning peak concurrent active users", "2,500 users (3-5x normal)", "Distributed spike testing simulating 8:00 - 10:00 AM traffic"),
        ("Reliability", "Maximum unplanned system downtime", "<= 43.8 minutes/month (99.9% SLA)", "Synthetic multi-region uptime probes (Pingdom/UptimeRobot)"),
        ("Availability", "Web platform and core booking API availability", ">= 99.9%", "APM uptime metrics and load balancer HTTP access logs"),
        ("Latency", "Read API response time (P95) for schedule searching", "< 300 ms", "Application Load Balancer access logs & APM tracing"),
        ("Security / HIPAA", "Data encryption in transit and at rest for PHI", "TLS 1.3 / AES-256-GCM", "Automated SAST/DAST audits, SSL Labs A+ check, KMS audit logs")
    ]
    while len(t_nfr.rows) > 1:
        t_nfr._tbl.remove(t_nfr.rows[-1]._tr)
    for row in nfr_data:
        r = t_nfr.add_row()
        for i in range(4):
            r.cells[i].text = row[i]

# Populate Golden Signals Table
if t_golden:
    golden_data = [
        ("Latency", "API Response Time P95", "Application Performance Monitoring (Datadog)", "< 500 ms", "High latency directly drives appointment abandonment and clinic frustration."),
        ("Traffic", "Active Concurrent Users & RPS", "Real-time user sessions & API Gateway RPS", "Baseline: 500 CCU / 150 RPS\nPeak: 2,500 CCU / 800 RPS", "Guides horizontal pod autoscaling before system saturation occurs."),
        ("Errors", "HTTP 5xx Server Error Rate", "APM error tracking & log aggregation (ELK)", "< 0.1% (Alert if > 0.5%)", "Failed booking transactions result in lost patient visits and revenue."),
        ("Saturation", "DB Connection Pool Utilization", "Database monitoring (PgBouncer metrics)", "< 70% (Alert if > 80%)", "Connection exhaustion causes cascading service crashes and dropped transactions.")
    ]
    while len(t_golden.rows) > 1:
        t_golden._tbl.remove(t_golden.rows[-1]._tr)
    for row in golden_data:
        r = t_golden.add_row()
        for i in range(5):
            r.cells[i].text = row[i]

# Populate Peak Traffic Preparation Table
if t_peak:
    peak_data = [
        ("Database Host CPU", "Alert if > 70% for 5 min", "Alert if > 85% for 5 min", "High read query volume and concurrent slot indexing consume CPU; avoids false alarm fatigue."),
        ("API Response Time (P95)", "Alert if > 500 ms", "Alert if > 1,200 ms", "Allows for minor queuing delay while remaining strictly under the 3.0s hard user SLA."),
        ("DB Connection Pool", "Alert if > 65%", "Alert if > 85%", "Horizontal pod autoscaling adds concurrent connections; buffered headroom prevents exhaustion."),
        ("Pod Memory Utilization", "Alert if > 75%", "Alert if > 85%", "Accommodates in-memory session caches during maximum concurrent active logins.")
    ]
    while len(t_peak.rows) > 1:
        t_peak._tbl.remove(t_peak.rows[-1]._tr)
    for row in peak_data:
        r = t_peak.add_row()
        for i in range(4):
            r.cells[i].text = row[i]

# Add explanation paragraph under Peak Traffic Preparation
doc.add_heading("Three-Layer Performance Optimization Plan", level=2)
doc.add_paragraph(
    "1. Infrastructure Layer:\n"
    "   - Scheduled Horizontal Pod Autoscaling (HPA): Proactively scale application pods from 5 to 20 replicas at 07:30 AM EST every Monday, pre-warming resources before traffic hits.\n"
    "   - Edge Caching & CDN: Cloudflare CDN offloads 95% of static frontend UI assets and compresses payloads via Brotli.\n\n"
    "2. Application Layer:\n"
    "   - In-Memory Redis Caching: Cache doctor availability slots with 15-second TTL, reducing database read queries by 75%. Cache invalidated on booking commit.\n"
    "   - Asynchronous Decoupling: RabbitMQ message queues decouple SMS dispatch and calendar sync, ensuring third-party API latency never blocks the booking flow.\n\n"
    "3. Database Layer:\n"
    "   - Composite Indexing: B-Tree indexes on (clinic_id, doctor_id, appointment_date, status) enable sub-10ms slot queries.\n"
    "   - Read-Write Splitting & Connection Pooling: PgBouncer caps connections at 300, routing search queries to read replicas and reserving Primary for atomic booking commits."
)

doc.add_heading("Circuit Breaking & Graceful Degradation Runbook", level=2)
doc.add_paragraph(
    "- Level 1 (Traffic > 2,500 CCU): Automatically defer non-essential background tasks (historical reports, non-urgent exports).\n"
    "- Level 2 (DB CPU > 88%): Throttle external calendar sync webhooks; queue updates for off-peak execution while keeping core booking live.\n"
    "- Level 3 (Emergency): Activate Cloudflare Virtual Waiting Room to queue incoming users and preserve core transaction stability."
)

doc.save(out2)
print(f"Successfully generated: {out2} ({os.path.getsize(out2)} bytes)")
